"""
Advibe Investor List ETL Pipeline
---------------------------------
Ingests, normalizes, deduplicates, and merges 'Investor list.xlsx'
into data/investors_raw.csv and data/people_raw.csv.
"""

import os
import re
import csv
import uuid
import openpyxl
from pathlib import Path
from urllib.parse import urlparse

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
EXCEL_PATH = REPO_ROOT / "Investor list.xlsx"
INV_CSV = REPO_ROOT / "data" / "investors_raw.csv"
PPL_CSV = REPO_ROOT / "data" / "people_raw.csv"

def clean(v):
    if v is None:
        return ""
    return str(v).strip()

def clean_url(u):
    if not u:
        return ""
    u = u.strip()
    if not u.startswith("http://") and not u.startswith("https://"):
        u = "https://" + u
    return u

def extract_domain(url, email):
    if url:
        u = clean_url(url)
        try:
            domain = urlparse(u).netloc
            if domain.startswith("www."):
                domain = domain[4:]
            if domain:
                return domain.lower()
        except:
            pass
    if email and "@" in email:
        domain = email.split("@")[-1].strip().lower()
        if "." in domain and not any(g in domain for g in ["gmail", "yahoo", "hotmail", "outlook", "icloud"]):
            return domain
    return ""

def domain_to_firm_name(domain):
    if not domain:
        return ""
    base = domain.split(".")[0]
    return base.capitalize()

DM_SIGNALS = [
    "partner", "managing partner", "general partner", "principal", 
    "founder", "co-founder", "director", "managing director", 
    "vp", "vice president", "head of", "investments", "investment",
    "cio", "chief investment", "president"
]

def is_decision_maker(title):
    if not title:
        return True
    t = title.lower()
    return any(sig in t for sig in DM_SIGNALS)

def infer_fund_type_and_track(firm_name, title, sector_notes=""):
    combo = f"{firm_name} {title} {sector_notes}".lower()
    if any(k in combo for k in ["real estate", "reit", "property", "multifamily", "housing", "land", "development"]):
        return "Real Estate PE / REIT", "real_estate"
    elif any(k in combo for k in ["family office", "endowment", "pension", "sovereign", "fund of funds", "allocator", "foundation", "superannuation"]):
        return "Family Office / LP Allocator", "fund_lp"
    elif any(k in combo for k in ["angel", "syndicate"]):
        return "Angel Syndicate", "venture"
    else:
        return "Venture Capital", "venture"

def run_etl():
    print(f"Loading {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    
    # 1. Read existing CSVs
    existing_investors = []
    existing_inv_by_name = {}
    if INV_CSV.exists():
        with open(INV_CSV, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                existing_investors.append(row)
                existing_inv_by_name[row["firm_name"].strip().lower()] = row
    
    existing_people = []
    existing_ppl_by_email = {}
    if PPL_CSV.exists():
        with open(PPL_CSV, "r", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                existing_people.append(row)
                if row.get("email"):
                    existing_ppl_by_email[row["email"].strip().lower()] = row

    print(f"Existing catalog: {len(existing_investors)} firms, {len(existing_people)} people.")

    # 2. Extract from Excel
    extracted_records = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [clean(c) for c in rows[0]]
        for r_idx, r in enumerate(rows[1:]):
            row_dict = {headers[i]: clean(r[i]) if i < len(r) else "" for i in range(len(headers))}
            
            full_name = ""
            role_title = ""
            company = ""
            phone = ""
            email = ""
            location = ""
            website = ""
            linkedin = ""
            
            if sname in ["Sheet1", "Sheet3"]:
                full_name = row_dict.get("Full Name", "")
                role_title = row_dict.get("Title", "")
                company = row_dict.get("Company", "")
                phone = row_dict.get("Phone Number", "")
                email = row_dict.get("Email", "")
                location = row_dict.get("Location", "")
                website = row_dict.get("Website", "")
                linkedin = row_dict.get("LinkedIn", "")
            elif sname == "Sheet2":
                company = row_dict.get("Firm Name", "")
                email = row_dict.get("Email ID", "")
                website = row_dict.get("Website/Application Link", "")
                role_title = "Investment Team"
                full_name = "" # Firm-level contact
            elif sname == "Sheet4":
                full_name = row_dict.get(headers[0], "")
                role_title = row_dict.get("Title", "")
                phone = row_dict.get("PhoneNumber", "")
                email = row_dict.get("Email", "")
                location = row_dict.get("Location", "")
                website = row_dict.get("Website", "")
                linkedin = row_dict.get("LinkedIn", "")
                dom = extract_domain(website, email)
                company = domain_to_firm_name(dom)
            
            if not company and not full_name:
                continue

            # Handle primary email
            primary_email = ""
            if email:
                split_emails = [e.strip() for e in re.split(r"[,;]+", email) if e.strip()]
                if split_emails:
                    primary_email = split_emails[0]

            extracted_records.append({
                "source": f"Investor_list.xlsx:{sname}",
                "full_name": full_name,
                "role_title": role_title or "Partner",
                "company": company or "Independent Capital",
                "phone": phone,
                "email": primary_email,
                "location": location,
                "website": clean_url(website),
                "linkedin": clean_url(linkedin)
            })

    print(f"Extracted {len(extracted_records)} raw records from Excel.")

    # 3. Deduplicate and collect firms & people
    new_firms_map = {} # firm_key -> firm_record
    deduped_people = []
    seen_ppl_keys = set(existing_ppl_by_email.keys())

    for rec in extracted_records:
        firm_raw = rec["company"].strip()
        if not firm_raw:
            continue
        firm_key = firm_raw.lower()
        
        # Determine firm ID
        if firm_key in existing_inv_by_name:
            inv_id = existing_inv_by_name[firm_key]["id"]
        elif firm_key in new_firms_map:
            inv_id = new_firms_map[firm_key]["id"]
        else:
            inv_id = f"a0000002-{len(new_firms_map):04d}-0000-0000-000000000001"
            fund_type, track = infer_fund_type_and_track(firm_raw, rec["role_title"])
            
            geo = rec["location"] if rec["location"] else "US|Global"
            new_firms_map[firm_key] = {
                "id": inv_id,
                "firm_name": firm_raw,
                "fund_type": fund_type,
                "aum": "500000000" if "lp" in track else "250000000",
                "stage_focus": "Seed|Series A|Series B" if track == "venture" else "Acquisition|Value-Add|Core",
                "sector_focus": "AI/ML|B2B SaaS|Fintech|Healthcare" if track == "venture" else "Multifamily|Commercial|Industrial" if track == "real_estate" else "Diversified Alternatives|Fund Commitments",
                "geography_focus": geo,
                "typical_check_min": "250000" if track == "venture" else "1000000",
                "typical_check_max": "5000000" if track == "venture" else "25000000",
                "website_url": rec["website"],
                "source": rec["source"]
            }

        # Handle Person
        p_name = rec["full_name"].strip() or f"{firm_raw} Partner"
        p_email = rec["email"].strip().lower()
        p_key = (p_email, firm_key) if p_email else (p_name.lower(), firm_key)

        if p_key in seen_ppl_keys:
            continue
        seen_ppl_keys.add(p_key)
        if p_email:
            seen_ppl_keys.add(p_email)

        p_id = f"b0000002-{len(deduped_people):04d}-0000-0000-000000000001"
        deduped_people.append({
            "id": p_id,
            "investor_id": inv_id,
            "full_name": p_name,
            "role_title": rec["role_title"] or "Partner",
            "email": rec["email"] or f"contact@{extract_domain(rec['website'], '') or 'fund.com'}",
            "linkedin_url": rec["linkedin"] or "",
            "is_decision_maker": "true" if is_decision_maker(rec["role_title"]) else "false",
            "verified": "true" if rec["email"] and "@" in rec["email"] else "false"
        })

    print(f"Generated {len(new_firms_map)} new unique firms and {len(deduped_people)} new unique people.")

    # 4. Write updated investors_raw.csv
    total_investors = existing_investors + list(new_firms_map.values())
    fieldnames_inv = [
        "id", "firm_name", "fund_type", "aum", "stage_focus", "sector_focus",
        "geography_focus", "typical_check_min", "typical_check_max", "website_url", "source"
    ]
    with open(INV_CSV, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames_inv)
        writer.writeheader()
        for inv in total_investors:
            writer.writerow({k: inv.get(k, "") for k in fieldnames_inv})

    # 5. Write updated people_raw.csv
    total_people = existing_people + deduped_people
    fieldnames_ppl = [
        "id", "investor_id", "full_name", "role_title", "email",
        "linkedin_url", "is_decision_maker", "verified"
    ]
    with open(PPL_CSV, "w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames_ppl)
        writer.writeheader()
        for ppl in total_people:
            writer.writerow({k: ppl.get(k, "") for k in fieldnames_ppl})

    print(f"\nSUCCESS: Catalog scaled to {len(total_investors)} total firms and {len(total_people)} total people.")

if __name__ == "__main__":
    run_etl()
